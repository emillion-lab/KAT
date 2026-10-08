#!/usr/bin/env python3
"""KAT v2 — „Трафик ↑ → ламарини ↑, но регистрираните не“: три теста, записани преди резултата.

Хипотеза (Емил, 08.10.2026): при силен трафик дребните удари растат, но се оправят с
двустранен протокол и не влизат в статистиката на МВР. Статистиката губи все по-голям дял.

Предсказания, които я подкрепят:
  T1 (години, страна): застрахователните претенции за имуществени вреди растат по-бързо
     от регистрираните от МВР ПТП с пострадали → съотношението претенции/ПТП расте.
  T2 (дни, София, 2021-01 → 2022-04): в дни с повече шофиране (Apple driving index)
     регистрираните ПТП НЕ растат пропорционално, а делът с пострадали ПАДА.
  T3 (месеци, София, 2024-01 → 2025-07): при по-голямо задръстване (TomTom) делът с пострадали пада.
Ако T1–T3 не излизат така — хипотезата не е подкрепена от тези данни.
"""
import json, math
import numpy as np, pandas as pd
import statsmodels.api as sm
import statsmodels.formula.api as smf
from scipy import stats
import openpyxl

RAW, OUT = "v2/raw", "v2"

# TomTom Traffic Index, София — месечно ниво на задръстване, % (tomtom.com/traffic-index/city/sofia, 08.10.2026).
# Графиката дава двойки (2025, 2024) по месеци; средните 46.5 и 47.1 съвпадат с публикуваните 47% и −0.5 п.п.
# От 08.2025 katastrofi.bg показва срив: ПТП с пострадали в София падат от ~60 на ~15 на месец.
# Това е смяна в източника, не реален спад — периодът се изключва и се маркира в отчета.
DATA_BREAK = "2025-08-01"

TOMTOM = {2025: [44, 46, 48, 50, 49, 42, 36, 33, 46, 57, 55, 52],
          2024: [44, 46, 48, 49, 48, 42, 36, 35, 48, 56, 58, 55]}


def t1():
    wb = openpyxl.load_workbook(f"{RAW}/kfn_ibnr_2023.xlsx", data_only=True)
    ws = wb["предявени-брой"]
    prop, inj, block = {}, {}, None
    for r in ws.iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        if not v:
            continue
        if isinstance(v[0], str) and "имуществени" in v[0]:
            block = inj if "неимуществени" in v[0] else prop
        elif isinstance(v[0], int) and 2000 < v[0] < 2100 and len(v) > 1 and block is not None:
            block[v[0]] = v[1]  # подадени в годината на събитието (лаг 0) — сравнимо за всички години
    expo = {}
    for r in wb["изложеност"].iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        if len(v) == 2 and isinstance(v[0], int) and v[0] not in expo:
            expo[v[0]] = v[1]
    mvr = {d["year"]: d for d in json.load(open("data/historical.json", encoding="utf-8"))["national"]}
    rows = []
    for y in sorted(prop):
        if y < 2015 or y not in mvr:  # 2015+: един и същ източник за МВР (до 2014 серията е друга)
            continue
        m = mvr[y]
        rows.append({"year": y, "claims_property": round(prop[y]), "claims_injury": round(inj.get(y, float("nan"))),
                     "mvr_ptp_casualties": m["total"], "mvr_injured": m["injured"], "vehicles": round(expo[y]),
                     "ratio": prop[y] / m["total"],
                     "claims_per_1000veh": 1000 * prop[y] / expo[y], "ptp_per_1000veh": 1000 * m["total"] / expo[y]})
    df = pd.DataFrame(rows)
    X = sm.add_constant(df["year"] - 2015)
    fit = sm.OLS(np.log(df["ratio"]), X).fit()
    pre = df[df.year <= 2019]["ratio"].mean()
    post = df[df.year >= 2021]["ratio"].mean()
    return df, {"trend_pct_per_year": round((math.exp(fit.params.iloc[1]) - 1) * 100, 2),
                "p": round(float(fit.pvalues.iloc[1]), 4), "ratio_2015_19": round(pre, 2), "ratio_2021_23": round(post, 2),
                "claims_growth_pct": round((df.claims_property.iloc[-1] / df.claims_property.iloc[0] - 1) * 100, 1),
                "ptp_growth_pct": round((df.mvr_ptp_casualties.iloc[-1] / df.mvr_ptp_casualties.iloc[0] - 1) * 100, 1),
                "veh_growth_pct": round((df.vehicles.iloc[-1] / df.vehicles.iloc[0] - 1) * 100, 1)}


def sofia_daily():
    d = pd.read_csv(f"{RAW}/sofia_accidents_daily.csv", parse_dates=["date"])
    full = pd.DataFrame({"date": pd.date_range(d.date.min(), d.date.max())})
    d = full.merge(d, on="date", how="left").fillna(0)
    d["dow"] = d.date.dt.dayofweek
    d["month"] = d.date.dt.month
    d["year"] = d.date.dt.year
    return d


def t2(d):
    a = pd.read_csv(f"{RAW}/apple_bg.csv", parse_dates=["date"])
    a = a[a["sub-region"] == "Total"][["date", "driving"]]
    m = d.merge(a, on="date").dropna(subset=["driving"])
    m = m[m.n > 0].copy()  # празни дни = липса в данните, не нула катастрофи
    m["drv_z"] = (m.driving - m.driving.mean()) / m.driving.std()
    nb = smf.negativebinomial("n ~ drv_z + C(dow) + C(month) + C(year)", data=m).fit(disp=0, maxiter=300)
    bi = smf.glm("with_casualties + I(n - with_casualties) ~ drv_z + C(dow) + C(month) + C(year)",
                 data=m, family=sm.families.Binomial()).fit()
    q = pd.qcut(m.driving, 4, labels=["най-малко", "под средното", "над средното", "най-много"])
    tab = m.groupby(q, observed=True).agg(days=("n", "size"), ptp_day=("n", "mean"),
                                          cas_share=("with_casualties", "sum"), tot=("n", "sum"))
    tab["cas_share"] = (100 * tab.cas_share / tab.tot).round(1)
    tab["ptp_day"] = tab.ptp_day.round(1)
    return {"days": int(len(m)), "from": str(m.date.min().date()), "to": str(m.date.max().date()),
            "count_irr": round(math.exp(nb.params["drv_z"]), 3), "count_p": round(float(nb.pvalues["drv_z"]), 4),
            "share_or": round(math.exp(bi.params["drv_z"]), 3), "share_p": round(float(bi.pvalues["drv_z"]), 4),
            "corr_raw": round(float(stats.spearmanr(m.driving, m.n).correlation), 3)}, tab.drop(columns="tot")


def t3(d):
    d = d[(d.year.isin([2024, 2025])) & (d.date < DATA_BREAK)]
    mo = d.groupby(["year", "month"]).agg(n=("n", "sum"), cas=("with_casualties", "sum")).reset_index()
    mo["congestion"] = [TOMTOM[y][m - 1] for y, m in zip(mo.year, mo.month)]
    mo["share"] = 100 * mo.cas / mo.n
    rn = stats.spearmanr(mo.congestion, mo.n)
    rs = stats.spearmanr(mo.congestion, mo.share)
    return {"months": int(len(mo)), "rho_count": round(float(rn.correlation), 3), "p_count": round(float(rn.pvalue), 4),
            "rho_share": round(float(rs.correlation), 3), "p_share": round(float(rs.pvalue), 4)}, mo


def yearly_sofia(d):
    y = d.groupby("year").agg(ptp=("n", "sum"), with_casualties=("with_casualties", "sum"),
                              days=("date", "size")).reset_index()
    y["share_pct"] = (100 * y.with_casualties / y.ptp).round(1)
    return y


def main():
    d = sofia_daily()
    r1, s1 = t1()
    s2, tab2 = t2(d)
    s3, mo = t3(d)
    ys = yearly_sofia(d)
    res = {"T1": s1, "T2": s2, "T3": s3}
    json.dump(res, open(f"{OUT}/link_results.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    r1.to_csv(f"{OUT}/link_t1_years.csv", index=False)
    mo.to_csv(f"{OUT}/link_t3_months.csv", index=False)

    L = ["# KAT v2 — трафик, ламарини и какво не влиза в статистиката", "",
         "Хипотеза: при силен трафик дребните удари растат, но се оправят с двустранен протокол и не влизат",
         "в статистиката на МВР. Три теста, записани в кода преди резултата.", "",
         "## T1 — Години, цялата страна: застрахователни претенции срещу ПТП на МВР", "",
         "Претенции = подадени в КФН за имуществени вреди по ГО в същата година (сравнимо за всички години).",
         "ПТП = регистрирани от МВР с пострадали (НСИ/МВР).", "",
         "| година | претенции (имущ.) | ПТП с пострадали (МВР) | претенции на 1 ПТП | застраховани МПС | претенции/1000 МПС | ПТП/1000 МПС |",
         "|---|---|---|---|---|---|---|"]
    for r in r1.itertuples():
        L.append(f"| {r.year} | {int(r.claims_property):,} | {int(r.mvr_ptp_casualties):,} | {r.ratio:.2f} | {int(r.vehicles):,} | "
                 f"{r.claims_per_1000veh:.1f} | {r.ptp_per_1000veh:.2f} |".replace(",", " "))
    L += ["", f"- Претенции 2015→2023: **{s1['claims_growth_pct']:+}%**; ПТП с пострадали: **{s1['ptp_growth_pct']:+}%**; "
              f"застраховани МПС: {s1['veh_growth_pct']:+}%.",
          f"- Претенции на 1 ПТП: средно **{s1['ratio_2015_19']}** (2015–19) → **{s1['ratio_2021_23']}** (2021–23); "
          f"тренд {s1['trend_pct_per_year']:+}% годишно, p={s1['p']}.", "",
          "## T2 — Дни, София: колко се кара срещу регистрирани ПТП и дял с пострадали", "",
          f"Apple Mobility „driving“ (България, дневно) × ПТП в София (katastrofi.bg/МВР), {s2['from']} → {s2['to']}, "
          f"{s2['days']} дни. Контроли: ден от седмицата, месец, година.", "",
          f"- Регистрирани ПТП на 1 SD повече шофиране: IRR **{s2['count_irr']}**, p={s2['count_p']}",
          f"- Шанс ПТП да е с пострадали на 1 SD повече шофиране: OR **{s2['share_or']}**, p={s2['share_p']}", "",
          "| шофиране (четвъртини) | дни | ПТП на ден | % с пострадали |", "|---|---|---|---|"]
    L += [f"| {k} | {int(r.days)} | {r.ptp_day} | {r.cas_share} |" for k, r in tab2.iterrows()]
    L += ["", "## T3 — Месеци, София 2024-01 → 2025-07: задръстване (TomTom) срещу ПТП", "",
          f"⚠ **Изключено: от {DATA_BREAK} нататък** — срив в данните (ПТП с пострадали падат ~4 пъти за един месец).", "",
          f"- Задръстване ↔ брой регистрирани ПТП: ρ = **{s3['rho_count']}**, p={s3['p_count']} ({s3['months']} месеца)",
          f"- Задръстване ↔ % с пострадали: ρ = **{s3['rho_share']}**, p={s3['p_share']}", "",
          "| месец | задръстване % | ПТП | % с пострадали |", "|---|---|---|---|"]
    L += [f"| {int(r.year)}-{int(r.month):02d} | {int(r.congestion)} | {int(r.n)} | {r.share:.1f} |" for _, r in mo.iterrows()]
    L += ["", "## София по години (katastrofi.bg)", "", "| година | ПТП | с пострадали | % | дни с данни |", "|---|---|---|---|---|"]
    L += [f"| {int(r.year)} | {int(r.ptp)} | {int(r.with_casualties)} | {r.share_pct}{' ⚠' if r.year >= 2025 else ''} | {int(r.days)} |" for _, r in ys.iterrows()]
    L += ["", f"⚠ 2025–2026: делът с пострадали е подбит от срива в данните от {DATA_BREAK}; не се ползва за изводи."]
    L += ["", "## Ограничения", "",
          "- КФН дава претенции за цялата страна, не за София, и само до 2023 г. (последният намерен файл).",
          "- Ръст на претенциите може да идва и от по-лесно подаване, измами или по-висока осведоменост, не само от повече удари.",
          "- Apple „driving“ е брой заявки за маршрут в Apple Maps за цялата страна — заместител на трафика, не броене на коли.",
          "- TomTom месечните стойности са от графиката на публичния индекс; 19 точки са малко.",
          "- T3 не може да отдели сезона: есента е и по-задръстена, и по-тъмна и мокра; лятото — по-свободно и по-бързо.",
          "- katastrofi.bg сменя източника през годините (2021–2023 от chernapista.com, после табло на МВР);",
          "  резкият спад на дела с пострадали през 2025–2026 трябва да се провери — може да е от данните, не от пътя.", ""]
    open(f"{OUT}/link_report.md", "w", encoding="utf-8").write("\n".join(L))
    print("\n".join(L))


if __name__ == "__main__":
    main()
